import csv
import io
import re
from typing import List, Dict, Any, Tuple
from fastapi import HTTPException
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from sqlalchemy import func
from models.maquina import Maquina


def normalizar_coluna(nome: str) -> str:
    """Mapeia variações de cabeçalhos comuns em planilhas para a chave correta."""
    texto = (nome or "").strip().lower()
    texto = (
        texto.replace("á", "a")
        .replace("ã", "a")
        .replace("â", "a")
        .replace("é", "e")
        .replace("ê", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ô", "o")
        .replace("õ", "o")
        .replace("ú", "u")
        .replace("ç", "c")
    )
    texto = re.sub(r"[^a-z0-9]", "", texto)

    if any(k in texto for k in ["datacoleta", "datadecoleta", "datacoletada", "datada", "data"]):
        return "data_coleta"
    if any(k in texto for k in ["nomedamaquina", "hostname", "maquina", "computador", "host", "equipamento"]):
        return "nome_maquina"
    if any(k in texto for k in ["usuarioad", "userad", "loginad", "aduser", "contaad"]) or texto == "ad":
        return "usuario_ad"
    if any(k in texto for k in ["usuario", "colaborador", "funcionario", "responsavel", "nome"]):
        return "usuario"
    if any(k in texto for k in ["outrosips", "outroip", "ipexterno", "ipsecundario", "outros"]):
        return "outros_ips"
    if any(k in texto for k in ["ipinterno", "enderecoip", "ipmaquina", "ip"]):
        return "ip"
    if any(k in texto for k in ["versaobuild", "buildoffice", "versaodooffice", "versao"]):
        return "office_versao"
    if any(k in texto for k in ["arquitetura", "arch"]):
        return "office_arquitetura"
    if any(k in texto for k in ["situacaooffice", "situacao"]):
        return "office_situacao"
    if any(k in texto for k in ["pacoteoffice", "office", "msoffice"]):
        return "office"
    if any(k in texto for k in ["sistemaoperacional", "sistema", "so", "windows", "os"]):
        return "sistema_operacional"
    if any(k in texto for k in ["departamento", "setor", "depto", "area"]):
        return "setor"
    if any(k in texto for k in ["antivirus", "anti-virus", "av", "edr"]):
        return "antivirus"
    if any(k in texto for k in ["observacoes", "observacao", "obs", "anotacoes", "nota"]):
        return "observacoes"
    if any(k in texto for k in ["statuscoleta", "statusdacoleta", "status"]):
        return "status"

    return texto


def simplificar_nome_office(texto: str | None) -> str | None:
    """Normaliza nomes longos do Office para formato enxuto (ex: Office 2013, Microsoft 365)."""
    if not texto:
        return None
    limpo = texto.strip()
    if not limpo or limpo.lower() in [
        "nao identificado",
        "não identificado",
        "falha na consulta",
        "falha",
        "sem office",
        "nenhum",
        "none",
        "null",
        "-",
        "—",
    ]:
        return None

    if re.search(r"\b(365|o365|m365)\b", limpo, re.IGNORECASE):
        return "Microsoft 365"

    m_ano = re.search(r"\b(20\d\d)\b", limpo)
    if m_ano:
        return f"Office {m_ano.group(1)}"

    return limpo


def parse_planilha_bytes(conteudo: bytes, filename: str) -> List[Dict[str, str]]:
    """Extrai lista de dicionários a partir do binário de um arquivo Excel ou CSV."""
    linhas_dados: List[Dict[str, str]] = []
    fname = filename.lower()

    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
            ws = wb.active
            todas_linhas = list(ws.iter_rows(values_only=True))
            if not todas_linhas:
                raise HTTPException(status_code=400, detail="A planilha está vazia.")

            cabecalhos_brutos = [str(c) if c is not None else "" for c in todas_linhas[0]]
            mapa_chaves = [normalizar_coluna(h) for h in cabecalhos_brutos]

            for row in todas_linhas[1:]:
                if not any(row):
                    continue
                item_dict = {}
                for idx, cell in enumerate(row):
                    if idx < len(mapa_chaves):
                        chave = mapa_chaves[idx]
                        if chave:
                            valor = str(cell).strip() if cell is not None else ""
                            item_dict[chave] = valor
                linhas_dados.append(item_dict)
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Erro ao processar planilha Excel: {str(e)}")

    elif fname.endswith(".csv") or fname.endswith(".txt"):
        texto = ""
        for encoding in ["utf-8", "latin-1", "cp1252"]:
            try:
                texto = conteudo.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not texto:
            raise HTTPException(status_code=400, detail="Não foi possível decodificar o arquivo CSV.")

        separador = ";" if ";" in texto.splitlines()[0] else ","
        reader = csv.reader(io.StringIO(texto), delimiter=separador)
        todas_linhas = list(reader)
        if not todas_linhas:
            raise HTTPException(status_code=400, detail="O arquivo CSV está vazio.")

        cabecalhos_brutos = todas_linhas[0]
        mapa_chaves = [normalizar_coluna(h) for h in cabecalhos_brutos]

        for row in todas_linhas[1:]:
            if not any(c.strip() for c in row if c):
                continue
            item_dict = {}
            for idx, cell in enumerate(row):
                if idx < len(mapa_chaves):
                    chave = mapa_chaves[idx]
                    if chave:
                        item_dict[chave] = cell.strip()
            linhas_dados.append(item_dict)
    else:
        raise HTTPException(
            status_code=400,
            detail="Formato não suportado. Envie um arquivo Excel (.xlsx) ou CSV (.csv).",
        )

    return linhas_dados


def processar_linhas_importadas(linhas_dados: List[Dict[str, str]], db: Session) -> Dict[str, Any]:
    """Salva ou atualiza as máquinas no banco."""
    criados = 0
    atualizados = 0
    ignorados = 0

    for linha in linhas_dados:
        nome_maquina = (linha.get("nome_maquina") or "").strip()
        if not nome_maquina:
            ignorados += 1
            continue

        ip = (linha.get("ip") or "").strip()
        usuario = (linha.get("usuario") or "").strip()
        usuario_ad = (linha.get("usuario_ad") or "").strip()
        if "\\" in usuario and not usuario_ad:
            partes = usuario.split("\\", 1)
            usuario_ad = partes[1].strip()
            usuario = partes[1].replace(".", " ").title()
        elif "\\" in usuario:
            partes = usuario.split("\\", 1)
            usuario = partes[1].replace(".", " ").title()

        # Limpar prefixo de domínio de rede (ex: ARTHICOM\controle.qualidade -> controle.qualidade)
        if "\\" in usuario_ad:
            usuario_ad = usuario_ad.split("\\", 1)[1].strip()
        if "/" in usuario_ad:
            usuario_ad = usuario_ad.split("/", 1)[1].strip()

        if usuario_ad.startswith("@"):
            usuario_ad = usuario_ad.lstrip("@").strip()

        setor = (linha.get("setor") or "").strip()
        if setor and setor.islower():
            setor = setor.capitalize()

        so = (linha.get("sistema_operacional") or "").strip()
        office_bruto = (linha.get("office") or "").strip()
        office = simplificar_nome_office(office_bruto) or ""
        av = (linha.get("antivirus") or "").strip()
        obs = (linha.get("observacoes") or "").strip()

        status_info = (linha.get("status") or "").strip()
        data_coleta = (linha.get("data_coleta") or "").strip()
        outros_ips = (linha.get("outros_ips") or "").strip()
        office_versao = (linha.get("office_versao") or "").strip()
        office_arquitetura = (linha.get("office_arquitetura") or "").strip()

        if not obs and (status_info or data_coleta or outros_ips or office_versao):
            detalhes = []
            if outros_ips:
                detalhes.append(f"Outro IP: {outros_ips}")
            if office_versao:
                arch_str = f" ({office_arquitetura})" if office_arquitetura and office_arquitetura != "Nao identificada" else ""
                detalhes.append(f"Build Office: {office_versao}{arch_str}")
            if status_info and status_info != "OK":
                detalhes.append(f"Status Coleta: {status_info}")
            if data_coleta:
                detalhes.append(f"Coletado via rede em: {data_coleta}")
            obs = " · ".join(detalhes)
        elif data_coleta and (not obs or "Coletado" in obs or "Status:" in obs):
            detalhes = []
            if outros_ips:
                detalhes.append(f"Outro IP: {outros_ips}")
            if office_versao:
                arch_str = f" ({office_arquitetura})" if office_arquitetura and office_arquitetura != "Nao identificada" else ""
                detalhes.append(f"Build Office: {office_versao}{arch_str}")
            if status_info and status_info != "OK":
                detalhes.append(f"Status Coleta: {status_info}")
            detalhes.append(f"Coletado via rede em: {data_coleta}")
            obs = " · ".join(detalhes)

        existente = (
            db.query(Maquina)
            .filter(
                func.lower(Maquina.nome_maquina) == nome_maquina.lower(),
                Maquina.ativo.is_(True),
            )
            .first()
        )

        if existente:
            if ip:
                existente.ip = ip
            if usuario:
                existente.usuario = usuario
            if usuario_ad:
                existente.usuario_ad = usuario_ad
            if setor:
                existente.setor = setor
            if so:
                existente.sistema_operacional = so
            if office:
                existente.office = office
            if av:
                existente.antivirus = av
            if obs:
                existente.observacoes = obs
            atualizados += 1
        else:
            nova = Maquina(
                usuario=usuario or None,
                nome_maquina=nome_maquina,
                usuario_ad=usuario_ad or None,
                ip=ip or None,
                sistema_operacional=so or None,
                office=office or None,
                setor=setor or None,
                antivirus=av or None,
                observacoes=obs or None,
                ativo=True,
            )
            db.add(nova)
            criados += 1

    db.commit()

    return {
        "mensagem": "Importação concluída com sucesso.",
        "total_linhas": len(linhas_dados),
        "criados": criados,
        "atualizados": atualizados,
        "ignorados": ignorados,
    }


def gerar_planilha_exportacao(db: Session) -> io.BytesIO:
    """Gera arquivo Excel formatado com todas as máquinas."""
    maquinas = (
        db.query(Maquina)
        .order_by(
            func.lower(Maquina.setor).asc().nullslast(),
            func.lower(Maquina.nome_maquina).asc(),
        )
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relação de IPs"

    headers = [
        "Nome da Máquina",
        "IP",
        "Usuário",
        "Usuário do AD",
        "Setor",
        "Sistema Operacional",
        "Office",
        "Antivírus",
        "Status",
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0033CC", end_color="0033CC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, m in enumerate(maquinas, start=2):
        ws.append(
            [
                m.nome_maquina or "",
                m.ip or "",
                m.usuario or "",
                m.usuario_ad or "",
                m.setor or "",
                m.sistema_operacional or "",
                m.office or "",
                m.antivirus or "",
                "Ativo" if m.ativo else "Inativo",
            ]
        )
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
